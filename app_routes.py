# 这个文件包含职称评审、培训管理和导出功能的路由
# 需要在app.py末尾导入这些路由

from flask import request, redirect, url_for, flash, render_template, send_file
from flask_login import login_required
from models import db, Employee, Title, Training
from datetime import datetime
from openpyxl import Workbook
from io import BytesIO

# ==================== 职称评审 ====================

def register_title_routes(app, admin_required):
    
    @app.route('/titles')
    @login_required
    def titles():
        """职称评审列表"""
        page = request.args.get('page', 1, type=int)
        employee_id = request.args.get('employee_id', type=int)
        status = request.args.get('status')
        
        query = Title.query
        
        # 筛选
        if employee_id:
            query = query.filter_by(employee_id=employee_id)
        if status:
            query = query.filter_by(status=status)
        
        # 分页
        pagination = query.order_by(Title.application_date.desc()).paginate(
            page=page, per_page=app.config['ITEMS_PER_PAGE'], error_out=False
        )
        
        employees = Employee.query.filter_by(status='在职').all()
        
        return render_template('titles/list.html',
                             titles=pagination.items,
                             pagination=pagination,
                             employees=employees,
                             employee_id=employee_id,
                             selected_status=status)


    @app.route('/titles/add', methods=['GET', 'POST'])
    @admin_required
    def add_title():
        """添加职称评审"""
        if request.method == 'POST':
            try:
                title = Title(
                    employee_id=request.form.get('employee_id'),
                    title_name=request.form.get('title_name'),
                    title_level=request.form.get('title_level'),
                    application_date=datetime.strptime(request.form.get('application_date'), '%Y-%m-%d').date() if request.form.get('application_date') else None,
                    approval_date=datetime.strptime(request.form.get('approval_date'), '%Y-%m-%d').date() if request.form.get('approval_date') else None,
                    status=request.form.get('status', '待审核'),
                    certificate_no=request.form.get('certificate_no'),
                    remarks=request.form.get('remarks')
                )
                
                db.session.add(title)
                db.session.commit()
                
                flash('职称评审添加成功！', 'success')
                return redirect(url_for('titles'))
            except Exception as e:
                db.session.rollback()
                flash(f'添加失败：{str(e)}', 'error')
        
        employees = Employee.query.filter_by(status='在职').all()
        return render_template('titles/add.html', employees=employees)


    @app.route('/titles/<int:id>/edit', methods=['GET', 'POST'])
    @admin_required
    def edit_title(id):
        """编辑职称评审"""
        title = Title.query.get_or_404(id)
        
        if request.method == 'POST':
            try:
                title.employee_id = request.form.get('employee_id')
                title.title_name = request.form.get('title_name')
                title.title_level = request.form.get('title_level')
                title.application_date = datetime.strptime(request.form.get('application_date'), '%Y-%m-%d').date() if request.form.get('application_date') else None
                title.approval_date = datetime.strptime(request.form.get('approval_date'), '%Y-%m-%d').date() if request.form.get('approval_date') else None
                title.status = request.form.get('status')
                title.certificate_no = request.form.get('certificate_no')
                title.remarks = request.form.get('remarks')
                
                db.session.commit()
                flash('职称评审更新成功！', 'success')
                return redirect(url_for('titles'))
            except Exception as e:
                db.session.rollback()
                flash(f'更新失败：{str(e)}', 'error')
        
        employees = Employee.query.filter_by(status='在职').all()
        return render_template('titles/edit.html', title=title, employees=employees)


    @app.route('/titles/<int:id>/delete', methods=['POST'])
    @admin_required
    def delete_title(id):
        """删除职称评审"""
        title = Title.query.get_or_404(id)
        try:
            db.session.delete(title)
            db.session.commit()
            flash('职称评审已删除', 'success')
        except Exception as e:
            db.session.rollback()
            flash(f'删除失败：{str(e)}', 'error')
        
        return redirect(url_for('titles'))


# ==================== 培训与发展 ====================

def register_training_routes(app, admin_required):
    
    @app.route('/trainings')
    @login_required
    def trainings():
        """培训列表"""
        page = request.args.get('page', 1, type=int)
        employee_id = request.args.get('employee_id', type=int)
        status = request.args.get('status')
        
        query = Training.query
        
        # 筛选
        if employee_id:
            query = query.filter_by(employee_id=employee_id)
        if status:
            query = query.filter_by(status=status)
        
        # 分页
        pagination = query.order_by(Training.start_date.desc()).paginate(
            page=page, per_page=app.config['ITEMS_PER_PAGE'], error_out=False
        )
        
        employees = Employee.query.filter_by(status='在职').all()
        
        return render_template('trainings/list.html',
                             trainings=pagination.items,
                             pagination=pagination,
                             employees=employees,
                             employee_id=employee_id,
                             selected_status=status)


    @app.route('/trainings/add', methods=['GET', 'POST'])
    @admin_required
    def add_training():
        """添加培训记录"""
        if request.method == 'POST':
            try:
                training = Training(
                    employee_id=request.form.get('employee_id'),
                    training_name=request.form.get('training_name'),
                    training_type=request.form.get('training_type'),
                    training_content=request.form.get('training_content'),
                    start_date=datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date() if request.form.get('start_date') else None,
                    end_date=datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None,
                    duration_hours=float(request.form.get('duration_hours', 0)) if request.form.get('duration_hours') else None,
                    location=request.form.get('location'),
                    instructor=request.form.get('instructor'),
                    cost=float(request.form.get('cost', 0)) if request.form.get('cost') else 0,
                    status=request.form.get('status', '已计划'),
                    completion_status=request.form.get('completion_status'),
                    remarks=request.form.get('remarks')
                )
                
                db.session.add(training)
                db.session.commit()
                
                flash('培训记录添加成功！', 'success')
                return redirect(url_for('trainings'))
            except Exception as e:
                db.session.rollback()
                flash(f'添加失败：{str(e)}', 'error')
        
        employees = Employee.query.filter_by(status='在职').all()
        return render_template('trainings/add.html', employees=employees)


    @app.route('/trainings/<int:id>/edit', methods=['GET', 'POST'])
    @admin_required
    def edit_training(id):
        """编辑培训记录"""
        training = Training.query.get_or_404(id)
        
        if request.method == 'POST':
            try:
                training.employee_id = request.form.get('employee_id')
                training.training_name = request.form.get('training_name')
                training.training_type = request.form.get('training_type')
                training.training_content = request.form.get('training_content')
                training.start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date() if request.form.get('start_date') else None
                training.end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
                training.duration_hours = float(request.form.get('duration_hours', 0)) if request.form.get('duration_hours') else None
                training.location = request.form.get('location')
                training.instructor = request.form.get('instructor')
                training.cost = float(request.form.get('cost', 0)) if request.form.get('cost') else 0
                training.status = request.form.get('status')
                training.completion_status = request.form.get('completion_status')
                training.remarks = request.form.get('remarks')
                
                db.session.commit()
                flash('培训记录更新成功！', 'success')
                return redirect(url_for('trainings'))
            except Exception as e:
                db.session.rollback()
                flash(f'更新失败：{str(e)}', 'error')
        
        employees = Employee.query.filter_by(status='在职').all()
        return render_template('trainings/edit.html', training=training, employees=employees)


    @app.route('/trainings/<int:id>/delete', methods=['POST'])
    @admin_required
    def delete_training(id):
        """删除培训记录"""
        training = Training.query.get_or_404(id)
        try:
            db.session.delete(training)
            db.session.commit()
            flash('培训记录已删除', 'success')
        except Exception as e:
            db.session.rollback()
            flash(f'删除失败：{str(e)}', 'error')
        
        return redirect(url_for('trainings'))


# ==================== Excel导出功能 ====================

def register_export_routes(app):
    
    @app.route('/export/employees')
    @login_required
    def export_employees():
        """导出员工信息到Excel"""
        employees = Employee.query.all()
        
        # 创建工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "员工信息"
        
        # 表头
        headers = ['工号', '姓名', '性别', '出生日期', '身份证号', '电话', '邮箱', 
                  '部门', '职位', '聘用类型', '入职日期', '状态', '学历', '专业', '毕业院校']
        ws.append(headers)
        
        # 数据
        for emp in employees:
            ws.append([
                emp.employee_no,
                emp.name,
                emp.gender,
                str(emp.birth_date) if emp.birth_date else '',
                emp.id_card,
                emp.phone,
                emp.email,
                emp.department.name if emp.department else '',
                emp.position,
                emp.employment_type,
                str(emp.hire_date) if emp.hire_date else '',
                emp.status,
                emp.education,
                emp.major,
                emp.graduate_school
            ])
        
        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'员工信息_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )


    @app.route('/export/attendances')
    @login_required
    def export_attendances():
        """导出考勤记录到Excel"""
        from models import Attendance
        attendances = Attendance.query.all()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "考勤记录"
        
        headers = ['员工姓名', '工号', '日期', '上班时间', '下班时间', '状态', '请假类型', '备注']
        ws.append(headers)
        
        for att in attendances:
            ws.append([
                att.employee.name,
                att.employee.employee_no,
                str(att.date),
                str(att.check_in_time) if att.check_in_time else '',
                str(att.check_out_time) if att.check_out_time else '',
                att.status,
                att.leave_type,
                att.remarks
            ])
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'考勤记录_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )


    @app.route('/export/performances')
    @login_required
    def export_performances():
        """导出绩效考核到Excel"""
        from models import Performance
        performances = Performance.query.all()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "绩效考核"
        
        headers = ['员工姓名', '工号', '年度', '季度', '教学评分', '科研评分', '服务评分', '总分', '等级', '考核人', '考核日期']
        ws.append(headers)
        
        for perf in performances:
            ws.append([
                perf.employee.name,
                perf.employee.employee_no,
                perf.year,
                perf.quarter,
                perf.teaching_score,
                perf.research_score,
                perf.service_score,
                perf.total_score,
                perf.rating,
                perf.evaluator,
                str(perf.evaluation_date) if perf.evaluation_date else ''
            ])
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'绩效考核_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )
